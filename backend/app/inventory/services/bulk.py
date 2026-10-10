"""Bulk price/quantity/discount update by Excel (T-39).

A supplier downloads a template (or exports their current products), edits the
sheet, and uploads it. Upload is two-phase: a *preview* validates every row and
colours it green/yellow/red (with a per-row reason) WITHOUT writing anything;
*apply* re-validates server-side and applies only the non-error rows. Discounts
go through the same best-price rule as the single-product screen (T-30).

Sheet columns (fixed order, row 1 = headers):
  A الباركود · B الاسم · C الكمية · D السعر · E الحد الأدنى · F الخصم
The discount cell is empty (no discount), "N%" (percent off), or a plain number
(the discounted unit price).
"""

from __future__ import annotations

import io
from decimal import Decimal, InvalidOperation
from typing import Any

from sqlalchemy import select

from ...extensions import db
from ...common.errors import ApiError, BadRequest
from ..models import Product
from . import offers as offers_svc

HEADERS = ["الباركود", "الاسم", "الكمية", "السعر", "الحد الأدنى", "الخصم"]
_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# Cap the rows we will parse so a crafted/huge (zip-bomb) sheet can't exhaust
# memory/CPU — a real bulk edit is hundreds of rows, not tens of thousands.
_MAX_ROWS = 5000


def _xlsx_safe(v: str) -> str:
    """Neutralise CSV/formula injection on EXPORT: a cell beginning with
    = + - @ (or a control char) is run as a formula by Excel/Sheets, so
    user-controlled text is prefixed with a single quote to stay data."""
    if v and v[0] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def _sheet(wb_active: Any) -> None:
    wb_active.sheet_view.rightToLeft = True
    wb_active.append(HEADERS)


def template_bytes() -> bytes:
    """An empty template with one illustrative example row."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "المنتجات"
    _sheet(ws)
    ws.append(["6221031492000", "مثال: صنف", 100, 50, 5, "10%"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_bytes(supplier_id: int) -> bytes:
    """The supplier's current products in the same shape, for edit + re-upload."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "المنتجات"
    _sheet(ws)
    for p in offers_svc.supplier_products(supplier_id):
        offer = p["offer"] or {}
        disc = ""
        if offer.get("discount_kind") == "percent" and offer.get("discount_value"):
            disc = f"{Decimal(offer['discount_value']).normalize()}%"
        elif offer.get("discount_kind") == "price" and offer.get("discount_value"):
            disc = str(Decimal(offer["discount_value"]).normalize())
        ws.append([
            _xlsx_safe(p["barcode"] or ""),
            _xlsx_safe(p["name"] or ""),
            float(Decimal(p["on_hand"])),
            float(Decimal(offer["unit_price"])) if offer.get("unit_price") else "",
            float(Decimal(offer["moq"])) if offer.get("moq") else 0,
            disc,
        ])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _num(cell: Any) -> Decimal | None:
    if cell is None or str(cell).strip() == "":
        return None
    return Decimal(str(cell).strip())


def _parse_discount(raw: Any) -> tuple[str, Decimal | None]:
    s = str(raw or "").strip()
    if not s:
        return ("none", None)
    if s.endswith("%"):
        return ("percent", Decimal(s[:-1].strip()))
    return ("price", Decimal(s))


def _validate_row(supplier_id: int, cells: list[Any]) -> dict[str, Any]:
    """Validate one row → {status, message, parsed}. status ∈ green|yellow|red."""
    barcode = str(cells[0] or "").strip() if len(cells) > 0 else ""
    name = str(cells[1] or "").strip() if len(cells) > 1 else ""
    row: dict[str, Any] = {"barcode": barcode, "name": name, "status": "red", "message": "", "parsed": None}

    if not barcode:
        row["message"] = "الباركود مفقود"
        return row
    product = db.session.execute(select(Product).where(Product.barcode == barcode)).scalar_one_or_none()
    if product is None:
        row["message"] = "منتج غير معروف بهذا الباركود"
        return row

    try:
        qty = _num(cells[2]) if len(cells) > 2 else None
        price = _num(cells[3]) if len(cells) > 3 else None
        moq = _num(cells[4]) if len(cells) > 4 else None
        kind, dvalue = _parse_discount(cells[5] if len(cells) > 5 else "")
    except (InvalidOperation, ValueError):
        row["message"] = "قيمة رقمية غير صالحة"
        return row

    if price is None or price <= 0:
        row["message"] = "السعر مطلوب وموجب"
        return row
    if qty is not None and qty < 0:
        row["message"] = "الكمية لا يمكن أن تكون سالبة"
        return row
    if moq is not None and moq < 0:
        row["message"] = "الحد الأدنى لا يمكن أن يكون سالبًا"
        return row

    notes: list[str] = []
    if kind != "none":
        if dvalue is None or dvalue <= 0:
            row["message"] = "قيمة الخصم غير صالحة"
            return row
        disc_price = offers_svc.discounted_price_for(unit_price=price, discount_kind=kind, discount_value=dvalue)
        try:
            offers_svc.validate_discount(product_id=product.id, supplier_id=supplier_id, discounted_price=disc_price)
        except ApiError as e:
            row["message"] = e.message  # the exact leak-free best-price message
            return row

    if qty is None:
        notes.append("الكمية فارغة — لن تتغيّر")

    row["parsed"] = {
        "product_id": product.id,
        "unit_price": str(price),
        "on_hand": str(qty) if qty is not None else None,
        "moq": str(moq) if moq is not None else "0",
        "discount_kind": kind,
        "discount_value": str(dvalue) if (kind != "none" and dvalue is not None) else None,
    }
    row["status"] = "yellow" if notes else "green"
    row["message"] = "؛ ".join(notes) if notes else "جاهز"
    return row


def _rows(data: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    assert ws is not None
    out: list[list[Any]] = []
    try:
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                continue  # header
            if row is None or all(c is None or str(c).strip() == "" for c in row):
                continue  # blank line
            # Only the first six columns matter; drop the rest so a sheet
            # padded with thousands of wide columns can't blow up memory.
            out.append(list(row[:6]))
            if len(out) > _MAX_ROWS:
                raise BadRequest(
                    f"الملف يتجاوز الحد الأقصى ({_MAX_ROWS} صف) — قسّمه إلى ملفات أصغر",
                    code="too_many_rows",
                )
    finally:
        wb.close()
    return out


def preview(*, supplier_id: int, data: bytes) -> dict[str, Any]:
    rows = [_validate_row(supplier_id, r) for r in _rows(data)]
    counts = {
        "green": sum(1 for r in rows if r["status"] == "green"),
        "yellow": sum(1 for r in rows if r["status"] == "yellow"),
        "red": sum(1 for r in rows if r["status"] == "red"),
    }
    # Preview only — never persist; roll back any read-side flushes.
    db.session.rollback()
    return {"rows": rows, "counts": counts}


def apply(*, supplier_id: int, data: bytes) -> dict[str, Any]:
    """Re-validate server-side and apply every non-red row. A row that fails at
    write time (e.g. a live price-lock) is flipped to red in the result."""
    results = [_validate_row(supplier_id, r) for r in _rows(data)]
    applied = 0
    for r in results:
        if r["status"] == "red" or r["parsed"] is None:
            continue
        p = r["parsed"]
        try:
            offers_svc.set_supplier_product(
                supplier_id=supplier_id,
                product_id=p["product_id"],
                unit_price=Decimal(p["unit_price"]),
                on_hand=Decimal(p["on_hand"]) if p["on_hand"] is not None else None,
                moq=Decimal(p["moq"]),
                is_active=True,
                discount_kind=p["discount_kind"],
                discount_value=Decimal(p["discount_value"]) if p["discount_value"] is not None else None,
            )
            applied += 1
        except ApiError as e:
            r["status"] = "red"
            r["message"] = e.message
        except Exception:  # pragma: no cover - defensive
            r["status"] = "red"
            r["message"] = "تعذّر الحفظ"
    return {"rows": results, "applied": applied, "total": len(results)}


def xlsx_response(data: bytes, filename: str):
    from flask import Response

    return Response(
        data, mimetype=_MIME,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
