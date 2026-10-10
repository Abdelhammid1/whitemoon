"""/catalog and /commerce blueprints — browse, cart, checkout, orders, RFQ."""

from __future__ import annotations

from decimal import Decimal
from html import escape
from typing import Any

from flask import Blueprint, Response, jsonify, request
from flask_jwt_extended import get_jwt_identity, jwt_required
from pydantic import ValidationError

from ..common.errors import ApiError, BadRequest, Forbidden, NotFound, Unauthorized
from ..extensions import db
from ..identity.models import User
from ..identity.services.audit import emit as audit_emit
from ..identity.services.rbac import has_permission, require_permission
from .models import Order
from .schemas import AddCartItemIn, CheckoutIn, RfqIn, RfqOfferIn, UpdateCartItemIn
from .services import cart as cart_svc
from .services import catalog as catalog_svc
from .services import orders as orders_svc
from .services import rfq as rfq_svc

bp = Blueprint("commerce", __name__)


def _parse(model_cls: Any) -> Any:
    try:
        return model_cls.model_validate(request.get_json(silent=True) or {})
    except ValidationError as e:
        raise BadRequest(f"Invalid payload: {e.errors()[0]['msg']}", code="validation_error") from e


def _uid() -> int:
    raw = get_jwt_identity()
    if raw is None:
        raise Unauthorized("No identity", code="no_identity")
    return int(raw)


def _kind(uid: int) -> str:
    user = db.session.get(User, uid)
    return user.kind if user else ""


def _xlsx_safe(v: str) -> str:
    """Neutralise CSV/formula injection: a cell beginning with = + - @ (or a
    control char) is executed as a formula by Excel/Sheets. User-controlled
    text (e.g. a store name) is prefixed with a single quote so it stays data."""
    if v and v[0] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


# ================================================================ catalog


@bp.get("/catalog/products")
@jwt_required()
def catalog_products():
    items = catalog_svc.browse(
        q=request.args.get("q"),
        category=request.args.get("category"),
        limit=min(int(request.args.get("limit", "100")), 500),
    )
    return jsonify({"items": items})


@bp.get("/catalog/products/<int:product_id>")
@jwt_required()
def catalog_product(product_id: int):
    item = catalog_svc.get_product(product_id)
    if item is None:
        raise NotFound("المنتج غير متاح", code="product_not_available")
    return jsonify(item)


@bp.get("/catalog/products/<int:product_id>/related")
@jwt_required()
def catalog_related(product_id: int):
    return jsonify({"items": catalog_svc.related_products(product_id)})


# ================================================================ cart


@bp.get("/commerce/cart")
@jwt_required()
def get_cart():
    return jsonify(cart_svc.serialize_cart(_uid()))


@bp.post("/commerce/cart/items")
@jwt_required()
def add_cart_item():
    payload = _parse(AddCartItemIn)
    item = cart_svc.add_item(customer_id=_uid(), offer_id=payload.offer_id, qty=Decimal(str(payload.qty)))
    return jsonify(cart_svc.serialize_cart(_uid())), 201 if item else 200


@bp.patch("/commerce/cart/items/<int:item_id>")
@jwt_required()
def update_cart_item(item_id: int):
    payload = _parse(UpdateCartItemIn)
    cart_svc.update_qty(customer_id=_uid(), item_id=item_id, qty=Decimal(str(payload.qty)))
    return jsonify(cart_svc.serialize_cart(_uid()))


@bp.delete("/commerce/cart/items/<int:item_id>")
@jwt_required()
def delete_cart_item(item_id: int):
    cart_svc.remove_item(customer_id=_uid(), item_id=item_id)
    return jsonify(cart_svc.serialize_cart(_uid()))


# ================================================================ checkout + orders


@bp.post("/commerce/checkout")
@jwt_required()
def checkout():
    payload = _parse(CheckoutIn)
    uid = _uid()
    order = orders_svc.checkout(
        customer_id=uid, payment_mode=payload.payment_mode, deferred_days=payload.deferred_days
    )
    audit_emit("commerce.order.placed", actor_user_id=uid, target_type="order", target_id=order.id)
    return jsonify(orders_svc.serialize_order(order, for_customer=True)), 201


@bp.post("/commerce/orders/<int:order_id>/reorder")
@jwt_required()
def reorder(order_id: int):
    """T-42: add a past order's items to the cart at current prices + warnings."""
    from .services import customer_shop as shop_svc

    return jsonify(shop_svc.reorder(customer_id=_uid(), order_id=order_id))


@bp.get("/commerce/usual-items")
@jwt_required()
def usual_items():
    """T-42: the customer's most-ordered items with usual qty + current price."""
    from .services import customer_shop as shop_svc

    return jsonify({"items": shop_svc.usual_items(customer_id=_uid())})


@bp.get("/commerce/usual-categories")
@jwt_required()
def usual_categories():
    """T-32: the customer's most-bought categories (last 90 days), else trending."""
    from .services import customer_shop as shop_svc

    return jsonify({"items": shop_svc.usual_categories(customer_id=_uid())})


@bp.get("/commerce/orders")
@jwt_required()
def list_orders():
    orders = orders_svc.list_orders(_uid())
    return jsonify({"items": [orders_svc.serialize_order(o, for_customer=True) for o in orders]})


@bp.get("/commerce/supplier/orders")
@jwt_required()
def supplier_orders():
    """The signed-in supplier's sub-orders with live status (T-06)."""
    uid = _uid()
    if not has_permission(uid, "offer.manage"):
        raise Forbidden("للموردين فقط", code="suppliers_only")
    return jsonify({"items": orders_svc.supplier_dashboard(uid)})


def _statement_access(customer_id: int) -> None:
    """Statement access: the customer, finance/admin, or an agent/branch inside
    the customer's geo territory (T-01/T-04)."""
    from ..partners.services import partners as partners_svc

    uid = _uid()
    allowed = (
        uid == customer_id
        or has_permission(uid, "user.read")
        or (partners_svc.is_partner(uid) and partners_svc.customer_in_scope(uid, customer_id))
    )
    if not allowed:
        raise Forbidden("خارج النطاق المصرح", code="out_of_scope")


def _parse_range() -> tuple[Any, Any]:
    from datetime import date as _date

    try:
        df = _date.fromisoformat(request.args["from"]) if request.args.get("from") else None
        dt = _date.fromisoformat(request.args["to"]) if request.args.get("to") else None
    except ValueError as e:
        raise BadRequest("صيغة التاريخ غير صحيحة (YYYY-MM-DD)", code="bad_date") from e
    return df, dt


def _build_statement(customer_id: int, date_from: Any = None, date_to: Any = None) -> dict[str, Any]:
    """Customer statement data (T-19): profile, orders, dues, approved payments,
    outstanding balance, credit limit and available credit. Optional date range
    filters orders (placed_at), dues (due_date) and payments (paid_on)."""
    from datetime import datetime, time

    from .models import Order
    from ..identity.models import CustomerProfile
    from ..sales.models import CustomerDue, PaymentApproval
    from ..sales.services import credit as credit_svc

    # The date range is applied in SQL (not fetch-then-filter), so a statement
    # for an old period returns exactly its rows; the cap is only a safety bound
    # that a date-bounded query is very unlikely to reach.
    _STMT_CAP = 5000
    dt_from = datetime.combine(date_from, time.min) if date_from else None
    dt_to = datetime.combine(date_to, time.max) if date_to else None

    def _range(col: Any, lo: Any, hi: Any, stmt: Any) -> Any:
        if lo is not None:
            stmt = stmt.where(col >= lo)
        if hi is not None:
            stmt = stmt.where(col <= hi)
        return stmt

    profile = db.session.get(CustomerProfile, customer_id)
    orders = list(
        db.session.execute(
            _range(
                Order.placed_at, dt_from, dt_to,
                db.select(Order).where(Order.customer_id == customer_id),
            ).order_by(Order.id.desc()).limit(_STMT_CAP)
        ).scalars()
    )
    dues = list(
        db.session.execute(
            _range(
                CustomerDue.due_date, date_from, date_to,
                db.select(CustomerDue).where(CustomerDue.customer_id == customer_id),
            ).order_by(CustomerDue.id.desc()).limit(_STMT_CAP)
        ).scalars()
    )
    payments = list(
        db.session.execute(
            _range(
                PaymentApproval.paid_on, date_from, date_to,
                db.select(PaymentApproval).where(
                    PaymentApproval.customer_id == customer_id,
                    PaymentApproval.status == "approved",
                ),
            ).order_by(PaymentApproval.id.desc()).limit(_STMT_CAP)
        ).scalars()
    )
    outstanding = credit_svc.outstanding(customer_id)
    limit = credit_svc.effective_limit(customer_id)
    tier = credit_svc.get_tier(customer_id)
    # T-28: batch-load the deferred-term snapshots (fee / annual % / days) for the
    # deferred orders so the statement can show the fee breakdown per order.
    from ..accounting.models import DeferredTerm

    deferred_order_ids = [o.id for o in orders if o.payment_mode == "deferred"]
    terms: dict[int, DeferredTerm] = {}
    if deferred_order_ids:
        terms = {
            t.order_id: t
            for t in db.session.execute(
                db.select(DeferredTerm).where(DeferredTerm.order_id.in_(deferred_order_ids))
            ).scalars()
        }
    # Light order rows (no sub-order/line lazy loads — the statement lists order
    # headers only, so this avoids an N+1 over a long order history).
    order_rows = []
    for o in orders:
        row: dict[str, Any] = {
            "id": o.id,
            "number": o.number,
            "status": o.status,
            "payment_mode": o.payment_mode,
            "total_cash": str(o.total_cash),
            "total_deferred": str(o.total_deferred),
            "placed_at": o.placed_at.isoformat(),
        }
        t = terms.get(o.id)
        if t is not None and t.fee is not None and t.annual_pct is not None:
            row["deferred_fee"] = str(t.fee)
            row["deferred_annual_pct"] = str(t.annual_pct)
            row["deferred_days"] = t.days
        order_rows.append(row)
    return {
        "customer_id": customer_id,
        "profile": {
            "display_name": profile.display_name if profile else None,
            "geo_area": profile.geo_area if profile else None,
        },
        "tier": tier.tier,
        "credit_limit": str(limit),
        "outstanding": str(outstanding),
        # The balance is the customer's CURRENT position (as of today), even when
        # a date range filters the rows below — labelled as such in the UI/PDF.
        "available": str(limit - outstanding),
        "as_of": datetime.now(credit_svc.BUSINESS_TZ).date().isoformat(),
        "orders": order_rows,
        "dues": [
            {
                "id": d.id,
                "amount": str(d.amount),
                "due_date": d.due_date.isoformat(),
                "status": d.status,
                "days_late": d.days_late,
            }
            for d in dues
        ],
        "payments": [
            {
                "id": p.id,
                "amount": str(p.amount),
                "paid_on": p.paid_on.isoformat(),
                "source": p.source,
            }
            for p in payments
        ],
    }


@bp.get("/commerce/customers/<int:customer_id>/statement")
@jwt_required()
def customer_statement(customer_id: int):
    """Unified customer statement: profile, orders, dues, payments, balance and
    available credit (T-04/T-19). Optional `from`/`to` date filter."""
    _statement_access(customer_id)
    df, dt = _parse_range()
    return jsonify(_build_statement(customer_id, df, dt))


@bp.get("/commerce/customers/<int:customer_id>/statement.pdf")
@jwt_required()
def customer_statement_pdf(customer_id: int):
    """Arabic PDF/printable statement (T-19)."""
    from ..common import documents as doc

    _statement_access(customer_id)
    df, dt = _parse_range()
    s = _build_statement(customer_id, df, dt)
    name = s["profile"]["display_name"] or f"#{customer_id}"

    due_rows = "".join(
        "<tr>"
        f"<td class='num'>{escape(d['due_date'])}</td>"
        f"<td class='num'>{doc.fmt_money(d['amount'])}</td>"
        f"<td>{_DUE_STATUS_AR.get(d['status'], d['status'])}</td>"
        "</tr>"
        for d in s["dues"]
    ) or "<tr><td colspan='3'>لا توجد ذمم</td></tr>"
    pay_rows = "".join(
        "<tr>"
        f"<td class='num'>{escape(p['paid_on'])}</td>"
        f"<td class='num'>{doc.fmt_money(p['amount'])}</td>"
        "</tr>"
        for p in s["payments"]
    ) or "<tr><td colspan='2'>لا توجد مدفوعات</td></tr>"

    rng = ""
    if df or dt:
        rng = f"<div><b>الفترة:</b> {escape(df.isoformat()) if df else '—'} → {escape(dt.isoformat()) if dt else '—'}</div>"
    body = (
        "<div class='head'>"
        "<div class='brand'>وايت مون<small>كشف حساب العميل</small></div>"
        f"<div class='meta'><div><b>العميل:</b> {escape(name)}</div>{rng}</div>"
        "</div>"
        "<h1>كشف الحساب</h1>"
        f"<p style='color:#6b7280;font-size:11px;margin:0 0 4px'>الرصيد والسقف محسوبان حتى تاريخه ({escape(s['as_of'])})، بصرف النظر عن فلتر الفترة.</p>"
        "<table class='totals'><tbody>"
        f"<tr><td>السقف الائتماني</td><td class='num'>{doc.fmt_money(s['credit_limit'])}</td></tr>"
        f"<tr><td>الرصيد المستحق</td><td class='num'>{doc.fmt_money(s['outstanding'])}</td></tr>"
        f"<tr><td>المتاح</td><td class='num'>{doc.fmt_money(s['available'])}</td></tr>"
        "</tbody></table>"
        "<h1 style='font-size:15px;margin-top:18px'>الذمم</h1>"
        "<table><thead><tr><th class='num'>تاريخ الاستحقاق</th><th class='num'>المبلغ</th><th>الحالة</th></tr></thead>"
        f"<tbody>{due_rows}</tbody></table>"
        "<h1 style='font-size:15px;margin-top:18px'>المدفوعات المعتمدة</h1>"
        "<table><thead><tr><th class='num'>التاريخ</th><th class='num'>المبلغ</th></tr></thead>"
        f"<tbody>{pay_rows}</tbody></table>"
        "<p style='color:#6b7280;font-size:11px'>جميع المبالغ بالجنيه المصري (ج.م).</p>"
    )
    html = doc.document(f"كشف حساب {name}", body)
    return doc.pdf_response(html, filename=f"statement-{customer_id}.pdf")


@bp.get("/commerce/customers/<int:customer_id>/statement.xlsx")
@jwt_required()
def customer_statement_xlsx(customer_id: int):
    """Excel (.xlsx) statement export (T-19)."""
    import io

    from openpyxl import Workbook

    _statement_access(customer_id)
    df, dt = _parse_range()
    s = _build_statement(customer_id, df, dt)
    name = s["profile"]["display_name"] or f"#{customer_id}"

    wb = Workbook()
    summary = wb.active
    assert summary is not None  # a new workbook always has an active sheet
    summary.title = "ملخص"
    summary.sheet_view.rightToLeft = True
    money = lambda v: float(Decimal(str(v)).quantize(Decimal("0.01")))  # noqa: E731 — 2dp, no binary drift
    summary.append(["العميل", _xlsx_safe(name)])
    summary.append(["السقف الائتماني", money(s["credit_limit"])])
    summary.append(["الرصيد المستحق", money(s["outstanding"])])
    summary.append(["المتاح", money(s["available"])])

    dues = wb.create_sheet("الذمم")
    dues.sheet_view.rightToLeft = True
    dues.append(["تاريخ الاستحقاق", "المبلغ", "الحالة", "أيام التأخير"])
    for d in s["dues"]:
        dues.append([d["due_date"], money(d["amount"]), _DUE_STATUS_AR.get(d["status"], d["status"]), d["days_late"]])

    pays = wb.create_sheet("المدفوعات")
    pays.sheet_view.rightToLeft = True
    pays.append(["التاريخ", "المبلغ"])
    for p in s["payments"]:
        pays.append([p["paid_on"], money(p["amount"])])

    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="statement-{customer_id}.xlsx"'},
    )


_DUE_STATUS_AR = {"open": "مفتوحة", "paid": "مسددة", "defaulted": "متعثرة", "cancelled": "ملغاة"}


@bp.get("/commerce/orders/<int:order_id>")
@jwt_required()
def get_order(order_id: int):
    order = db.session.get(Order, order_id)
    if order is None:
        raise NotFound("Order not found", code="order_not_found")
    uid = _uid()
    is_admin = has_permission(uid, "user.read")  # admin/staff
    if not is_admin and order.customer_id != uid:
        raise Forbidden("ليس طلبك", code="forbidden")
    return jsonify(orders_svc.serialize_order(order, for_customer=not is_admin))


_ORDER_STATUS_AR = {
    "pending": "قيد الانتظار",
    "confirmed": "مؤكَّد",
    "fulfilled": "منفَّذ",
    "cancelled": "ملغى",
}
_PAYMENT_MODE_AR = {"cash": "نقدي", "deferred": "آجل", "mixed": "مختلط"}


@bp.get("/commerce/orders/<int:order_id>/invoice")
@jwt_required()
def order_invoice(order_id: int):
    """Arabic PDF invoice for an order (T-23). Customer-facing: suppliers are
    never shown — lines are merged. The order's customer or staff may open it.

    Returns a PDF where WeasyPrint is available (production), otherwise a
    printable HTML page the browser can save as PDF."""
    from ..common import documents as doc
    from ..inventory.models import Product

    order = db.session.get(Order, order_id)
    if order is None:
        raise NotFound("الطلب غير موجود", code="order_not_found")
    uid = _uid()
    is_staff = has_permission(uid, "user.read")
    if not is_staff and order.customer_id != uid:
        raise Forbidden("ليس طلبك", code="forbidden")

    data = orders_svc.serialize_order(order, for_customer=True)

    # T-28: for a deferred order, show the transparent fee breakdown from the
    # per-order snapshot (annual %, duration, computed fee) + the due date.
    deferred_line = ""
    if order.payment_mode == "deferred":
        from ..accounting.models import DeferredTerm
        from ..sales.models import CustomerDue

        term = db.session.execute(
            db.select(DeferredTerm).where(DeferredTerm.order_id == order.id)
        ).scalar_one_or_none()
        due = db.session.execute(
            db.select(CustomerDue).where(CustomerDue.order_id == order.id)
        ).scalar_one_or_none()
        if term is not None and term.fee is not None and term.annual_pct is not None:
            due_txt = f" — تستحق في {due.due_date.isoformat()}" if due is not None else ""
            deferred_line = (
                f"<tr><td>رسوم الأجل ({Decimal(term.annual_pct):g}% سنويًا × "
                f"{term.days} يوم ÷ 365){escape(due_txt)}</td>"
                f"<td class='num'>{doc.fmt_money(term.fee)}</td></tr>"
            )

    product_ids = {ln["product_id"] for ln in data["lines"]}
    names = {
        p.id: p.name_ar
        for p in db.session.execute(
            db.select(Product).where(Product.id.in_(product_ids or {0}))
        ).scalars()
    }

    rows = "".join(
        "<tr>"
        f"<td>{escape(names.get(ln['product_id'], '—'))}</td>"
        f"<td class='num'>{doc.fmt_money(ln['unit_price'])}</td>"
        f"<td class='num'>{Decimal(ln['qty']):g}</td>"
        f"<td class='num'>{doc.fmt_money(ln['line_total'])}</td>"
        "</tr>"
        for ln in data["lines"]
    ) or "<tr><td colspan='4'>لا توجد بنود</td></tr>"

    is_deferred = order.payment_mode == "deferred"
    if is_deferred:
        totals_rows = (
            f"<tr><td>السعر النقدي</td><td class='num'>{doc.fmt_money(data['total_cash'])}</td></tr>"
            f"{deferred_line}"
        )
        grand_total = Decimal(data["total_deferred"])
        total_label = "الإجمالي الآجل"
    else:
        totals_rows = ""
        grand_total = Decimal(data["total_cash"])
        total_label = "الإجمالي"
    body = (
        "<div class='head'>"
        "<div class='brand'>وايت مون<small>فاتورة طلب</small></div>"
        "<div class='meta'>"
        f"<div><b>رقم الفاتورة:</b> {escape(data['number'])}</div>"
        f"<div><b>التاريخ:</b> {escape(data['placed_at'][:10])}</div>"
        f"<div><b>الحالة:</b> <span class='chip'>{_ORDER_STATUS_AR.get(data['status'], data['status'])}</span></div>"
        f"<div><b>طريقة الدفع:</b> {_PAYMENT_MODE_AR.get(data['payment_mode'], data['payment_mode'])}</div>"
        "</div></div>"
        "<h1>فاتورة</h1>"
        "<table><thead><tr>"
        "<th>الصنف</th><th class='num'>سعر الوحدة</th>"
        "<th class='num'>الكمية</th><th class='num'>الإجمالي</th>"
        "</tr></thead><tbody>"
        f"{rows}</tbody></table>"
        "<table class='totals'><tbody>"
        f"{totals_rows}"
        "</tbody><tfoot>"
        f"<tr><td>{total_label}</td><td class='num'>{doc.fmt_money(grand_total)}</td></tr>"
        "</tfoot></table>"
        "<p style='color:#6b7280;font-size:11px'>جميع المبالغ بالجنيه المصري (ج.م). "
        "هذه فاتورة مبدئية؛ الفاتورة الضريبية الإلكترونية (ETA) تصدر في مرحلة لاحقة.</p>"
    )
    html = doc.document(f"فاتورة {data['number']}", body)
    return doc.pdf_response(html, filename=f"invoice-{data['number']}.pdf")


@bp.get("/commerce/admin/orders")
@require_permission("order.manage")
def admin_list_orders():
    """All orders with search + status/customer/date filters (T-13)."""
    cid = request.args.get("customer_id", type=int)
    items = orders_svc.list_all_orders(
        status=request.args.get("status"),
        customer_id=cid,
        q=request.args.get("q"),
        date_from=request.args.get("date_from"),
        date_to=request.args.get("date_to"),
        limit=min(request.args.get("limit", 200, type=int), 500),
    )
    return jsonify({"items": orders_svc.serialize_admin_rows(items)})


@bp.post("/commerce/orders/<int:order_id>/<action>")
@require_permission("order.manage")
def transition_order(order_id: int, action: str):
    """confirm / fulfill / cancel — propagates to sub-orders; cancel also
    reverses the order's financials (T-13). The service validates the action."""
    order = orders_svc.transition_order(order_id, action, actor_id=_uid())
    audit_emit(
        f"commerce.order.{action}", actor_user_id=_uid(),
        target_type="order", target_id=order.id,
    )
    return jsonify(orders_svc.serialize_admin(order))


# ================================================================ RFQ


@bp.post("/commerce/rfqs")
@jwt_required()
def create_rfq():
    payload = _parse(RfqIn)
    uid = _uid()
    kind = _kind(uid)
    initiator_type = "supplier" if kind == "supplier" else "customer"
    rfq = rfq_svc.create_rfq(
        initiator_user_id=uid,
        initiator_type=initiator_type,
        product_id=payload.product_id,
        qty=Decimal(str(payload.qty)),
        deadline=payload.deadline,
        qualification_requirements=payload.qualification_requirements,
    )
    audit_emit("commerce.rfq.create", actor_user_id=uid, target_type="rfq", target_id=rfq.id)
    return jsonify(rfq_svc.serialize_rfq(rfq)), 201


@bp.get("/commerce/rfqs")
@jwt_required()
def list_rfqs():
    """RFQ list: a supplier's open inbox, a customer's own RFQs, or all (admin).

    The supplier inbox never reveals who opened each RFQ (mediated)."""
    uid = _uid()
    is_admin = has_permission(uid, "user.read")
    items = rfq_svc.list_rfqs(requester_id=uid, kind=_kind(uid), is_admin=is_admin)
    return jsonify({"items": items})


@bp.get("/commerce/rfqs/<int:rfq_id>")
@jwt_required()
def get_rfq(rfq_id: int):
    from .models import Rfq

    rfq = db.session.get(Rfq, rfq_id)
    if rfq is None:
        raise NotFound("RFQ not found", code="rfq_not_found")
    return jsonify(rfq_svc.serialize_rfq(rfq))


@bp.post("/commerce/rfqs/<int:rfq_id>/offers")
@jwt_required()
def submit_rfq_offer(rfq_id: int):
    uid = _uid()
    if _kind(uid) != "supplier":
        raise Forbidden("المورد فقط يقدّم عروض RFQ", code="supplier_only")
    payload = _parse(RfqOfferIn)
    offer = rfq_svc.submit_offer(
        rfq_id=rfq_id, supplier_id=uid, unit_price=Decimal(str(payload.unit_price)), moq=Decimal(str(payload.moq))
    )
    audit_emit("commerce.rfq.offer", actor_user_id=uid, target_type="rfq_offer", target_id=offer.id)
    return jsonify({"offer_id": offer.id, "rfq_id": rfq_id}), 201


@bp.get("/commerce/rfqs/<int:rfq_id>/offers")
@jwt_required()
def list_rfq_offers(rfq_id: int):
    uid = _uid()
    is_admin = has_permission(uid, "user.read")
    offers = rfq_svc.list_offers(rfq_id=rfq_id, requester_id=uid, is_admin=is_admin)
    return jsonify({"items": offers})


# ================================================================ errors


@bp.app_errorhandler(ApiError)
def handle_api_error(err: ApiError):
    return jsonify({"error": err.code, "message": err.message}), err.status
